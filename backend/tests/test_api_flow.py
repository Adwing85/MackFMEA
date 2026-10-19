"""End-to-End-Workflow über die REST-API: PDM-Import → Klassifikation → FMEA → Review → Export → PDM."""

import io

from openpyxl import load_workbook

from .conftest import login

CLASSIFICATION = {
    "function": "Last tragen",
    "sc_level": 3,
    "values": {"BA": "SCHW", "GEO": "PRIS", "MAT": "S355", "BEL": "SCHWELL", "LSP": "N3", "INS": "GUT", "SD": ["ERM", "RISS"]},
    "justification": "Primäres Tragwerk",
}


def _import_hubwerk(client, headers):
    response = client.post("/api/pdm/import", json={"drawing_number": "BG-300100", "function": "Last heben"}, headers=headers)
    assert response.status_code == 201, response.text
    return response.json()


def _classify_all(client, headers, assembly, sc_level=3):
    for component in assembly["components"]:
        body = {**CLASSIFICATION, "sc_level": sc_level}
        response = client.post(f"/api/components/{component['id']}/classifications", json=body, headers=headers)
        assert response.status_code == 201, response.text


def _generate(client, headers, assembly_id):
    response = client.post(f"/api/assemblies/{assembly_id}/fmea/generate", json={"team": "Konstruktion, QS"}, headers=headers)
    assert response.status_code == 201, response.text
    return response.json()


def test_login_rejects_wrong_password(client):
    assert client.post("/api/auth/login", json={"username": "admin", "password": "falsch"}).status_code == 401
    assert client.get("/api/assemblies").status_code == 401


def test_full_workflow(client, designer, reviewer, admin):
    assembly = _import_hubwerk(client, designer)
    assert assembly["source"] == "pdm:demo"
    assert len(assembly["components"]) == 5
    _classify_all(client, designer, assembly)

    doc = _generate(client, designer, assembly["id"])
    assert doc["status"] == "entwurf"
    assert doc["version"] == 1
    assert doc["rows"] and doc["requirements"]
    assert all(r["special_characteristic"] == "SC3" for r in doc["rows"])
    assert {s["step"] for s in doc["snapshot"]["process"]} == {1, 2, 3, 4, 5, 6}
    assert doc["outdated_components"] == []

    # Prüfer darf Entwürfe nicht bearbeiten, Konstrukteur schon
    row = doc["rows"][0]
    assert client.put(f"/api/fmea/{doc['id']}/rows/{row['id']}", json={"remarks": "x"}, headers=reviewer).status_code == 403
    edited = client.put(
        f"/api/fmea/{doc['id']}/rows/{row['id']}",
        json={"occurrence": 9, "detection": 9, "change_reason": "Test"},
        headers=designer,
    ).json()
    assert edited["action_priority"] == "H"
    assert edited["manually_edited"] is True

    assert client.post(f"/api/fmea/{doc['id']}/submit", headers=designer).json()["status"] == "in_pruefung"

    # Konstrukteur kann nicht freigeben; AP-H-Zeilen ohne Verantwortlichen blockieren die Freigabe
    assert client.post(f"/api/fmea/{doc['id']}/approve", json={}, headers=designer).status_code == 403
    blocked = client.post(f"/api/fmea/{doc['id']}/approve", json={}, headers=reviewer)
    assert blocked.status_code == 422
    current = client.get(f"/api/fmea/{doc['id']}", headers=reviewer).json()
    for r in current["rows"]:
        if r["action_priority"] == "H":
            response = client.put(
                f"/api/fmea/{doc['id']}/rows/{r['id']}", json={"opt_responsible": "K. Konstrukteur"}, headers=reviewer
            )
            assert response.status_code == 200
    approved = client.post(f"/api/fmea/{doc['id']}/approve", json={"comment": "ok"}, headers=reviewer).json()
    assert approved["status"] == "freigegeben"
    assert approved["reviewed_by"] == "pruefer"

    # Freigegebene FMEA ist schreibgeschützt
    assert client.put(f"/api/fmea/{doc['id']}/rows/{row['id']}", json={"remarks": "x"}, headers=admin).status_code == 409

    # Export
    xlsx = client.get(f"/api/fmea/{doc['id']}/export.xlsx", headers=designer)
    assert xlsx.status_code == 200
    wb = load_workbook(io.BytesIO(xlsx.content))
    assert wb.sheetnames == ["DFMEA", "Anforderungen", "Klassifikation", "Rückverfolgbarkeit"]
    assert wb["DFMEA"]["B8"].value == "1. Nächsthöhere Ebene"
    pdf = client.get(f"/api/fmea/{doc['id']}/export.pdf", headers=designer)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")

    # Rückspeicherung ins PDM
    stored = client.post(f"/api/fmea/{doc['id']}/pdm-store", headers=designer).json()
    assert stored["pdm_status"] == "gespeichert"
    assert "demo-pdm://" in stored["pdm_ref"]

    # Klassifikationsänderung → neue Version, alte FMEA wird als veraltet erkannt
    component = assembly["components"][0]
    response = client.post(
        f"/api/components/{component['id']}/classifications",
        json={**CLASSIFICATION, "sc_level": 4, "justification": "Neubewertung"},
        headers=designer,
    )
    assert response.json()["version"] == 2
    assert client.get(f"/api/fmea/{doc['id']}", headers=designer).json()["outdated_components"]

    doc2 = _generate(client, designer, assembly["id"])
    assert doc2["version"] == 2
    diff = client.get(f"/api/fmea/{doc2['id']}/compare/{doc['id']}", headers=designer).json()
    assert diff["rows_changed"] or diff["rows_added"]
    assert diff["requirements_added"]

    # Neue Freigabe ersetzt die alte Version
    for r in doc2["rows"]:
        if r["action_priority"] == "H":
            client.put(f"/api/fmea/{doc2['id']}/rows/{r['id']}", json={"opt_responsible": "QS"}, headers=designer)
    client.post(f"/api/fmea/{doc2['id']}/submit", headers=designer)
    assert client.post(f"/api/fmea/{doc2['id']}/approve", json={}, headers=reviewer).json()["status"] == "freigegeben"
    assert client.get(f"/api/fmea/{doc['id']}", headers=designer).json()["status"] == "ersetzt"

    history = client.get(f"/api/fmea/{doc['id']}/history", headers=reviewer).json()
    actions = {h["action"] for h in history}
    assert {"GENERATE", "SUBMIT", "APPROVE", "EXPORT", "PDM_STORE", "SUPERSEDE"} <= actions
    assert client.get("/api/audit/verify", headers=admin).json()["valid"] is True


def test_four_eyes_principle(client, admin, reviewer):
    assembly = _import_hubwerk(client, admin)
    _classify_all(client, admin, assembly, sc_level=0)
    doc = _generate(client, admin, assembly["id"])
    client.post(f"/api/fmea/{doc['id']}/submit", headers=admin)
    response = client.post(f"/api/fmea/{doc['id']}/approve", json={}, headers=admin)
    assert response.status_code == 403
    assert "Vier-Augen" in response.json()["detail"]
    rejected = client.post(f"/api/fmea/{doc['id']}/reject", json={"comment": "Team fehlt"}, headers=reviewer).json()
    assert rejected["status"] == "abgelehnt"
    assert client.post(f"/api/fmea/{doc['id']}/submit", headers=admin).json()["status"] == "in_pruefung"


def test_classification_requires_justification_from_sc2(client, designer):
    assembly = _import_hubwerk(client, designer)
    component_id = assembly["components"][0]["id"]
    body = {**CLASSIFICATION, "justification": ""}
    response = client.post(f"/api/components/{component_id}/classifications", json=body, headers=designer)
    assert response.status_code == 422
    body = {**CLASSIFICATION, "sc_level": 1, "justification": ""}
    assert client.post(f"/api/components/{component_id}/classifications", json=body, headers=designer).status_code == 201


def test_invalid_classification_is_rejected(client, designer):
    assembly = _import_hubwerk(client, designer)
    component_id = assembly["components"][0]["id"]
    body = {**CLASSIFICATION, "values": {**CLASSIFICATION["values"], "MAT": "HOLZ"}}
    response = client.post(f"/api/components/{component_id}/classifications", json=body, headers=designer)
    assert response.status_code == 422
    assert any("Material" in e for e in response.json()["detail"])


def test_generate_without_classification_fails(client, designer):
    assembly = _import_hubwerk(client, designer)
    response = client.post(f"/api/assemblies/{assembly['id']}/fmea/generate", json={}, headers=designer)
    assert response.status_code == 422


def test_csv_import_endpoint(client, designer):
    csv = b"Pos;Zeichnungsnummer;Benennung;Menge\n10;X-1;Platte;1\n20;X-2;Bolzen;4\n"
    response = client.post(
        "/api/assemblies/import",
        files={"file": ("X-100.csv", csv, "text/csv")},
        data={"name": "Testgruppe"},
        headers=designer,
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["drawing_number"] == "X-100"
    assert body["name"] == "Testgruppe"
    assert [c["drawing_number"] for c in body["components"]] == ["X-1", "X-2"]


def test_role_permissions_on_catalog(client, designer, reviewer, admin):
    sc = {"name": "Neu"}
    assert client.put("/api/catalog/sc-classes/2", json=sc, headers=designer).status_code == 403
    assert client.put("/api/catalog/sc-classes/2", json=sc, headers=reviewer).status_code == 403
    assert client.put("/api/catalog/sc-classes/2", json=sc, headers=admin).status_code == 200
    reviewer_assembly = client.post("/api/pdm/import", json={"drawing_number": "BG-300100"}, headers=reviewer)
    assert reviewer_assembly.status_code == 403


def test_admin_adds_option_and_keys_are_generated(client, admin):
    attributes = client.get("/api/catalog/attributes", headers=admin).json()
    ba = next(a for a in attributes if a.get("code") == "BA")
    option = client.post(
        f"/api/catalog/attributes/{ba['id']}/options",
        json={
            "code": "ADD",
            "label": "Additiv gefertigt",
            "failure_cause": "Porosität / Bindefehler",
            "relevant_damages": ["ERM"],
        },
        headers=admin,
    )
    assert option.status_code == 201, option.text
    keys = client.get(f"/api/catalog/keys?option_id={option.json()['id']}", headers=admin).json()
    assert [k["code"] for k in keys] == [f"BA-ADD-SC{level}" for level in range(5)]

    measures = client.get("/api/catalog/measures", headers=admin).json()
    ut = next(m for m in measures if m["code"] == "M-E-08")
    key_sc2 = keys[2]
    response = client.post(
        f"/api/catalog/keys/{key_sc2['id']}/measures",
        json={"measure_id": ut["id"], "extent": "100 %", "propagate_up": True},
        headers=admin,
    )
    assert response.status_code == 200
    keys = client.get(f"/api/catalog/keys?option_id={option.json()['id']}", headers=admin).json()
    assert [len(k["measures"]) for k in keys] == [0, 0, 1, 1, 1]
    assert client.get("/api/catalog/consistency", headers=admin).json()["consistent"] is True

    bad = client.post(
        f"/api/catalog/attributes/{ba['id']}/options",
        json={"code": "X1", "label": "x", "relevant_damages": ["GIBTSNICHT"]},
        headers=admin,
    )
    assert bad.status_code == 422


def test_requirement_condition_validation(client, admin):
    response = client.post(
        "/api/catalog/requirements",
        json={"code": "R-T-1", "category": "Prüfmerkmal", "text": "x", "conditions": {"BA": ["HOLZ"]}},
        headers=admin,
    )
    assert response.status_code == 422
    response = client.post(
        "/api/catalog/requirements",
        json={"code": "R-T-1", "category": "Prüfmerkmal", "text": "x", "sc_min": 3, "sc_max": 1},
        headers=admin,
    )
    assert response.status_code == 422


def test_user_management(client, admin):
    created = client.post(
        "/api/users",
        json={"username": "neu", "full_name": "Neue Person", "role": "pruefer", "password": "geheim1234"},
        headers=admin,
    )
    assert created.status_code == 201
    login(client, "neu", "geheim1234")
    me = client.get("/api/auth/me", headers=admin).json()
    assert client.put(f"/api/users/{me['id']}", json={"role": "pruefer"}, headers=admin).status_code == 400


def test_dashboard(client, designer):
    data = client.get("/api/dashboard", headers=designer).json()
    assert data["assemblies"] == 1
    assert data["classified"] == 6
    assert sum(data["sc_distribution"].values()) == 6
