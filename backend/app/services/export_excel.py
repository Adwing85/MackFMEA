"""Excel-Export des FMEA-Dokuments im Formblatt-Aufbau nach AIAG & VDA (Schritte 1–6)."""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ..config import settings
from ..models import FmeaDocument
from .rules import RuleSet

STATUS_LABELS = {
    "entwurf": "Entwurf",
    "in_pruefung": "In Prüfung",
    "freigegeben": "Freigegeben",
    "abgelehnt": "Abgelehnt",
    "ersetzt": "Ersetzt",
}
AP_FILL = {"H": "F8B4B4", "M": "FCE7A2", "L": "BBE5C3"}
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
GROUP_FILLS = {
    "struktur": "DCE6F1",
    "funktion": "E2EFDA",
    "fehler": "FDE9D9",
    "risiko": "FFF2CC",
    "optimierung": "EDE7F6",
}
THIN = Side(style="thin", color="9CA3AF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

# (Gruppe, Überschrift, Breite, Attribut / Funktion)
COLUMNS: list[tuple[str, str, int, Any]] = [
    ("", "Nr.", 5, None),
    ("struktur", "1. Nächsthöhere Ebene", 22, "higher_level"),
    ("struktur", "2. Betrachtetes Element", 24, "focus_element"),
    ("struktur", "3. Nächstniedrigere Ebene / Merkmal", 24, "lower_level"),
    ("funktion", "1. Funktion nächsthöhere Ebene", 26, "function_higher"),
    ("funktion", "2. Funktion betrachtetes Element", 26, "function_focus"),
    ("funktion", "3. Funktion / Anforderung nächstniedrigere Ebene", 24, "function_lower"),
    ("fehler", "1. Fehlerfolge (FE)", 30, "failure_effect"),
    ("fehler", "B", 5, "severity"),
    ("fehler", "2. Fehlerart (FM)", 20, "failure_mode"),
    ("fehler", "3. Fehlerursache (FC)", 30, "failure_cause"),
    ("risiko", "Aktuelle Vermeidungsmaßnahmen (PC)", 42, "prevention_controls"),
    ("risiko", "A", 5, "occurrence"),
    ("risiko", "Aktuelle Entdeckungsmaßnahmen (DC)", 42, "detection_controls"),
    ("risiko", "E", 5, "detection"),
    ("risiko", "AP", 5, "action_priority"),
    ("risiko", "Besondere Merkmale", 9, "special_characteristic"),
    ("risiko", "RPZ (informativ)", 8, lambda r: r.severity * r.occurrence * r.detection),
    ("risiko", "Schlüssel", 18, lambda r: "\n".join(r.key_codes or [])),
    ("optimierung", "Vermeidungsmaßnahme", 34, "opt_prevention"),
    ("optimierung", "Entdeckungsmaßnahme", 34, "opt_detection"),
    ("optimierung", "Verantwortlich", 14, "opt_responsible"),
    ("optimierung", "Zieltermin", 11, "opt_target_date"),
    ("optimierung", "Status", 11, "opt_status"),
    ("optimierung", "A (neu)", 6, "opt_occurrence"),
    ("optimierung", "E (neu)", 6, "opt_detection_rating"),
    ("optimierung", "AP (neu)", 6, "opt_action_priority"),
    ("optimierung", "Bemerkungen", 30, "remarks"),
]
GROUP_TITLES = {
    "struktur": "Strukturanalyse (Schritt 2)",
    "funktion": "Funktionsanalyse (Schritt 3)",
    "fehler": "Fehleranalyse (Schritt 4)",
    "risiko": "Risikoanalyse (Schritt 5)",
    "optimierung": "Optimierung (Schritt 6)",
}


def _fmt_dt(value: Any) -> str:
    return value.strftime("%d.%m.%Y %H:%M") if value else "–"


def _header_block(ws: Worksheet, doc: FmeaDocument) -> int:
    last_col = len(COLUMNS)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    title = ws.cell(row=1, column=1, value=f"Design-FMEA (DFMEA) nach AIAG & VDA – {doc.title}")
    title.font = Font(size=14, bold=True, color="FFFFFF")
    title.fill = HEADER_FILL
    title.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 24

    assembly = doc.assembly
    fields = [
        [
            ("Unternehmen", settings.company_name or "–"),
            ("Betrachtungsgegenstand", f"{assembly.drawing_number} {assembly.name} (Rev. {assembly.revision or '–'})"),
            ("Konstruktionsverantwortung", doc.responsible or "–"),
            ("Vertraulichkeitsstufe", doc.confidentiality or "–"),
        ],
        [
            ("DFMEA-Nr.", doc.fmea_number),
            ("Version", str(doc.version)),
            ("Status", STATUS_LABELS.get(doc.status, doc.status)),
            ("Erstellt", f"{_fmt_dt(doc.created_at)} ({doc.created_by.username if doc.created_by else '–'})"),
        ],
        [
            ("Funktionsübergreifendes Team", doc.team or "–"),
            (
                "Geprüft / freigegeben",
                f"{doc.reviewed_by.username if doc.reviewed_by else '–'} {_fmt_dt(doc.reviewed_at) if doc.reviewed_at else ''}".strip(),
            ),
            ("Prüfkommentar", doc.review_comment or "–"),
            ("Inhalts-Hash (SHA-256)", doc.content_hash[:32] + "…" if doc.content_hash else "–"),
        ],
    ]
    block_cols = [(1, 3, 4, 8), (9, 11, 12, 15), (16, 19, 20, last_col)]
    for block, (lbl_start, lbl_end, val_start, val_end) in zip(fields, block_cols, strict=True):
        for offset, (label, value) in enumerate(block):
            row = 2 + offset
            ws.merge_cells(start_row=row, start_column=lbl_start, end_row=row, end_column=lbl_end)
            ws.merge_cells(start_row=row, start_column=val_start, end_row=row, end_column=val_end)
            lc = ws.cell(row=row, column=lbl_start, value=label)
            lc.font = Font(bold=True, size=9)
            lc.fill = PatternFill("solid", fgColor="F3F4F6")
            vc = ws.cell(row=row, column=val_start, value=value)
            vc.font = Font(size=9)
            vc.alignment = Alignment(wrap_text=True, vertical="center")
    return 7


def _rows_sheet(ws: Worksheet, doc: FmeaDocument) -> None:
    ws.title = "DFMEA"
    group_row = _header_block(ws, doc)
    head_row = group_row + 1

    # Gruppenüberschriften
    col = 1
    while col <= len(COLUMNS):
        group = COLUMNS[col - 1][0]
        end = col
        while end < len(COLUMNS) and COLUMNS[end][0] == group:
            end += 1
        if group:
            ws.merge_cells(start_row=group_row, start_column=col, end_row=group_row, end_column=end)
            cell = ws.cell(row=group_row, column=col, value=GROUP_TITLES[group])
            cell.font = Font(bold=True)
            cell.alignment = CENTER
            cell.fill = PatternFill("solid", fgColor=GROUP_FILLS[group])
        col = end + 1

    for idx, (group, title, width, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=head_row, column=idx, value=title)
        cell.font = Font(bold=True, size=9)
        cell.alignment = CENTER
        cell.border = BORDER
        if group:
            cell.fill = PatternFill("solid", fgColor=GROUP_FILLS[group])
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[head_row].height = 42

    for number, row in enumerate(doc.rows, start=1):
        r = head_row + number
        for idx, (_, title, _, accessor) in enumerate(COLUMNS, start=1):
            if accessor is None:
                value: Any = number
            elif callable(accessor):
                value = accessor(row)
            else:
                value = getattr(row, accessor)
            cell = ws.cell(row=r, column=idx, value=value if value not in (None, "") else None)
            cell.border = BORDER
            cell.font = Font(size=9)
            cell.alignment = (
                CENTER
                if title
                in {"Nr.", "B", "A", "E", "AP", "RPZ (informativ)", "A (neu)", "E (neu)", "AP (neu)", "Besondere Merkmale"}
                else WRAP
            )
            if title in {"AP", "AP (neu)"} and value in AP_FILL:
                cell.fill = PatternFill("solid", fgColor=AP_FILL[value])
                cell.font = Font(size=9, bold=True)
    ws.freeze_panes = ws.cell(row=head_row + 1, column=4)
    ws.auto_filter.ref = f"A{head_row}:{get_column_letter(len(COLUMNS))}{head_row + max(len(doc.rows), 1)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{group_row}:{head_row}"


def _table(ws: Worksheet, headers: list[tuple[str, int]], rows: list[list[Any]], start_row: int = 1) -> None:
    for idx, (title, width) in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=idx, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(idx)].width = width
    for r_idx, values in enumerate(rows, start=start_row + 1):
        for c_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.alignment = WRAP
            cell.border = BORDER
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    if rows:
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{start_row + len(rows)}"


def build_workbook(doc: FmeaDocument, ruleset: RuleSet) -> bytes:
    wb = Workbook()
    _rows_sheet(wb.active, doc)

    ws = wb.create_sheet("Anforderungen")
    scope_label = {"halbzeug": "Halbzeug", "bauteil": "Bauteil"}
    _table(
        ws,
        [("Bauteil", 32), ("SC", 6), ("Bereich", 11), ("Kategorie", 22), ("Code", 11), ("Norm", 13), ("Anforderung", 90)],
        [
            [
                r.component_label,
                f"SC{r.sc_level}",
                scope_label.get(r.scope, r.scope),
                r.category,
                r.requirement_code,
                r.standard_code,
                r.text,
            ]
            for r in doc.requirements
        ],
    )

    ws = wb.create_sheet("Klassifikation")
    attrs = ruleset.attributes
    headers = [("Pos.", 7), ("Zeichnungsnr.", 14), ("Bezeichnung", 26), ("Bezeichnung 2", 22), ("Funktion", 36), ("SC", 6)]
    headers += [(f"{a.name} (Liste {a.list_no})", 24) for a in attrs]
    headers += [("Kl.-Version", 9), ("Schlüssel", 40), ("Maßnahmen", 40), ("Anforderungen", 30)]
    rows = []
    for comp in (doc.snapshot or {}).get("components", []):
        values = {**comp.get("liste1", {}), **comp.get("liste2", {})}
        line = [
            comp.get("position"),
            comp.get("drawing_number"),
            values.get("bezeichnung"),
            values.get("bezeichnung2"),
            values.get("funktion"),
            f"SC{comp.get('sc_level')}",
        ]
        for a in attrs:
            codes = values.get(a.code) or []
            labels = []
            for code in codes:
                option = ruleset.option(a.code, code)
                labels.append(option.label if option else code)
            line.append(", ".join(labels))
        line += [
            comp.get("classification_version"),
            "\n".join(comp.get("keys", [])),
            ", ".join(comp.get("measures", [])),
            ", ".join(comp.get("requirements", [])),
        ]
        rows.append(line)
    _table(ws, headers, rows)

    ws = wb.create_sheet("Rückverfolgbarkeit")
    snapshot = doc.snapshot or {}
    info = [
        ["FMEA-Nummer", doc.fmea_number],
        ["Version", doc.version],
        ["Status", STATUS_LABELS.get(doc.status, doc.status)],
        ["Generiert am", snapshot.get("generated_at", "")],
        ["Generiert von", snapshot.get("generated_by", "")],
        ["Generator-Version", snapshot.get("generator_version", "")],
        ["Regelwerksstand (Fingerprint)", snapshot.get("ruleset_fingerprint", "")],
        ["Inhalts-Hash", doc.content_hash],
        ["Quelle Baugruppe", (snapshot.get("assembly") or {}).get("source", "")],
        ["PDM-Referenz", doc.pdm_ref or (snapshot.get("assembly") or {}).get("pdm_ref", "")],
    ]
    info += [[f"Schritt {p['step']}: {p['name']}", p["result"]] for p in snapshot.get("process", [])]
    info += [["Hinweis", w] for w in (doc.warnings or [])]
    _table(ws, [("Merkmal", 42), ("Wert", 100)], info)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
